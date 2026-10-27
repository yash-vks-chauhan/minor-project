from __future__ import annotations

import csv
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


BASE_DIR = Path(__file__).resolve().parents[1]
OUTPUT_DIR = BASE_DIR / "output" / "doc"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

DOCX_PATH = OUTPUT_DIR / "minor_project_agile_report.docx"
MD_PATH = OUTPUT_DIR / "minor_project_agile_report.md"
CSV_PATH = OUTPUT_DIR / "ms_planner_backlog.csv"
SCRUM_CSV_PATH = OUTPUT_DIR / "daily_scrum_updates.csv"

DATA_SOURCE_EVIDENCE = [
    "TRAI - Indian Telecom Services Performance Indicators Reports portal (trai.gov.in)",
    "RBI - Handbook of Statistics on Indian Economy portal (rbi.org.in)",
    "NITI Aayog - SDG India Index and Dashboard portal (sdgindiaindex.niti.gov.in)",
    "NPCI Product Statistics and Dataful Insights state-wise UPI analysis (npci.org.in, insights.dataful.in)",
    "CPCB - Annual Report on Implementation of E-Waste Management Rules",
    "Census of India 2011 - Primary Census Abstract (censusindia.gov.in)",
]


def count_files(path: Path) -> int:
    return sum(1 for item in path.iterdir() if item.is_file())


def build_epics() -> list[dict]:
    return [
        {
            "id": "E1",
            "objective": "Build a verified multi-source state dataset for the study.",
            "epic": "Data Acquisition and Verification",
            "priority": "Very High",
            "status": "Completed",
            "implementation": (
                "Integrated official and researcher-provided data from TRAI, RBI, "
                "NITI Aayog, NPCI/Dataful, CPCB, and Census, then merged only "
                "complete state records into the final analytical dataset."
            ),
            "evidence": [
                "src/01_data_engineering.py",
                *DATA_SOURCE_EVIDENCE,
                "README.md - data-source verification section",
                "RESEARCH_FINDINGS.md - methodology and source documentation",
            ],
            "user_stories": [
                {
                    "id": "US1",
                    "title": "Collect verified source data",
                    "story": (
                        "As a researcher, I want all required state-level indicators "
                        "collected from official or directly traceable sources so that "
                        "the study uses trustworthy inputs."
                    ),
                    "acceptance": [
                        "Raw source files are available for telecom, income, SDG, UPI, e-waste, and demographic indicators.",
                        "Each source is mapped to a state-wise field used in the study.",
                        "Data provenance is documented in the project report and repository.",
                    ],
                },
                {
                    "id": "US2",
                    "title": "Merge only complete state records",
                    "story": (
                        "As an analyst, I want the final dataset to contain only "
                        "complete state records so that no hypothesis test depends on "
                        "imputed or fabricated values."
                    ),
                    "acceptance": [
                        "The merge strategy uses an inner join on the State field.",
                        "The final unified dataset contains 16 states with complete values.",
                        "No missing-value imputation is applied to the primary dataset.",
                    ],
                },
            ],
        },
        {
            "id": "E2",
            "objective": "Engineer composite indicators and visualize state-level patterns.",
            "epic": "Feature Engineering and Exploratory Analysis",
            "priority": "Very High",
            "status": "Completed",
            "implementation": (
                "Computed Digital Access Score, CE Behavior Score, DCI, and derived "
                "features such as e-waste per capita, then generated publication-ready "
                "EDA figures for trends, distributions, rankings, and correlations."
            ),
            "evidence": [
                "src/01_data_engineering.py",
                "src/02_eda_visualizations.py",
                *DATA_SOURCE_EVIDENCE,
                "RESEARCH_FINDINGS.md - composite index and methodology documentation",
                "figures/fig1_correlation_heatmap.jpg",
                "figures/fig2_digital_vs_ce_scatter.jpg",
                "figures/fig3_dci_ranking.jpg",
                "figures/fig6_upi_trends.jpg",
            ],
            "user_stories": [
                {
                    "id": "US3",
                    "title": "Compute reproducible composite scores",
                    "story": (
                        "As a researcher, I want Digital Access Score, CE Behavior "
                        "Score, and DCI to be computed with documented formulas so "
                        "states can be compared consistently."
                    ),
                    "acceptance": [
                        "Derived fields include EWaste_PerCapita, Digital_Access_Score, CE_Behavior_Score, and DCI.",
                        "Weights and formulas are documented in the report and paper.",
                        "Descriptive statistics are exported for the final dataset.",
                    ],
                },
                {
                    "id": "US4",
                    "title": "Create visual summaries of the data",
                    "story": (
                        "As a supervisor, I want visual summaries of distributions, "
                        "trends, and state rankings so that the project findings are "
                        "easy to interpret."
                    ),
                    "acceptance": [
                        "EDA figures are generated for correlation, scatter, ranking, distributions, and UPI trends.",
                        "Figure files are saved in the project figures directory.",
                        "Visual outputs are suitable for report or paper inclusion.",
                    ],
                },
            ],
        },
        {
            "id": "E3",
            "objective": "Test the main research hypotheses with statistical methods.",
            "epic": "Statistical Hypothesis Testing",
            "priority": "High",
            "status": "Completed",
            "implementation": (
                "Ran OLS, correlation tests, multiple regression, mediation analysis, "
                "threshold regression, and robustness diagnostics to evaluate the "
                "relationship between digital inclusion and circular economy behavior."
            ),
            "evidence": [
                "src/03_statistical_analysis.py",
                "results/ols_primary_results.txt",
                "results/multiple_regression_results.txt",
                "results/mediation_analysis.txt",
                "results/threshold_analysis.txt",
                "results/robustness_checks.txt",
            ],
            "user_stories": [
                {
                    "id": "US5",
                    "title": "Test the primary hypothesis",
                    "story": (
                        "As a research team member, I want the central relationship "
                        "between Digital Access Score and CE Behavior Score tested "
                        "statistically so that the study has defensible evidence."
                    ),
                    "acceptance": [
                        "Pearson, Spearman, and Kendall correlations are computed.",
                        "An OLS model reports coefficient estimates, p-value, and R-squared.",
                        "Results are saved as reusable report artifacts.",
                    ],
                },
                {
                    "id": "US6",
                    "title": "Check alternative explanations",
                    "story": (
                        "As a reviewer, I want controlled regression, mediation, and "
                        "threshold tests so that the project checks competing "
                        "interpretations of the main result."
                    ),
                    "acceptance": [
                        "A multiple regression with controls is executed and documented.",
                        "Mediation and piecewise threshold analyses are generated.",
                        "Diagnostic tests for heteroskedasticity, autocorrelation, and influence are reported.",
                    ],
                },
            ],
        },
        {
            "id": "E4",
            "objective": "Address reviewer concerns and strengthen methodological defense.",
            "epic": "Reviewer Defense and Robustness Validation",
            "priority": "High",
            "status": "Completed",
            "implementation": (
                "Prepared additional analyses to handle small sample size, weighting "
                "sensitivity, alternative mediation paths, robust regression, and "
                "extended-sample checks for defense readiness."
            ),
            "evidence": [
                "src/04_reviewer_defense.py",
                "results/reviewer_defense_report.txt",
                "results/weight_sensitivity.csv",
                "figures/fig10_bootstrap_ci.jpg",
                "figures/fig11_weight_sensitivity.jpg",
                "figures/fig12_extended_sample.jpg",
            ],
            "user_stories": [
                {
                    "id": "US7",
                    "title": "Quantify robustness for small-sample results",
                    "story": (
                        "As a project evaluator, I want bootstrap confidence intervals "
                        "and sensitivity checks so that the conclusions remain credible "
                        "despite a small sample."
                    ),
                    "acceptance": [
                        "Bootstrap confidence intervals are generated for the main relationship.",
                        "Alternative weighting schemes are tested and summarized.",
                        "Outputs are stored as result files and figures for viva or review use.",
                    ],
                },
                {
                    "id": "US8",
                    "title": "Prepare defense-ready supporting analyses",
                    "story": (
                        "As a reviewer, I want alternative mediation paths and "
                        "extended-sample checks so that the study can answer likely "
                        "critical questions during evaluation."
                    ),
                    "acceptance": [
                        "Reviewer-defense scripts generate a dedicated defense report.",
                        "Extended-sample or alternate-path analyses are documented.",
                        "The report explains how each additional check addresses a review risk.",
                    ],
                },
            ],
        },
        {
            "id": "E5",
            "objective": "Validate predictive capability using comparative machine learning models.",
            "epic": "Predictive Modeling Validation",
            "priority": "High",
            "status": "Completed",
            "implementation": (
                "Compared seven machine-learning models using leave-one-out "
                "cross-validation across digital-only and full-feature sets, then "
                "exported prediction metrics and state-level predictions."
            ),
            "evidence": [
                "src/05_predictive_modeling.py",
                "results/ml_model_comparison_digital_only.csv",
                "results/ml_model_comparison_full.csv",
                "results/ml_state_predictions.csv",
                "results/predictive_modeling_report.txt",
                "figures/fig13_model_comparison.jpg",
                "figures/fig14_actual_vs_predicted.jpg",
                "figures/fig15_prediction_errors.jpg",
            ],
            "user_stories": [
                {
                    "id": "US9",
                    "title": "Compare multiple predictive models",
                    "story": (
                        "As a researcher, I want multiple machine-learning models "
                        "evaluated under leave-one-out validation so that predictive "
                        "claims are not dependent on a single method."
                    ),
                    "acceptance": [
                        "At least seven models are evaluated using leave-one-out cross-validation.",
                        "Both digital-only and full-feature model sets are compared.",
                        "Comparison metrics are exported in report-ready files.",
                    ],
                },
                {
                    "id": "US10",
                    "title": "Explain predictive performance clearly",
                    "story": (
                        "As a faculty evaluator, I want prediction plots and "
                        "state-level outputs so that machine-learning performance is "
                        "easy to interpret."
                    ),
                    "acceptance": [
                        "Actual-vs-predicted and error visualizations are generated.",
                        "State-level prediction outputs are exported.",
                        "A predictive modeling report summarizes the best model and metrics.",
                    ],
                },
            ],
        },
        {
            "id": "E6",
            "objective": "Produce advanced explainability, clustering, and final documentation outputs.",
            "epic": "Advanced ML, Explainability, and Reporting",
            "priority": "High",
            "status": "Completed",
            "implementation": (
                "Executed XGBoost, nested cross-validation, SHAP explainability, "
                "clustering, PCA, and prepared final documentation artifacts "
                "including the paper and research findings."
            ),
            "evidence": [
                "src/06_advanced_ml.py",
                "results/nested_cv_results.csv",
                "results/shap_values_per_state.csv",
                "results/state_clusters.csv",
                "results/advanced_ml_report.txt",
                "paper.tex",
                "RESEARCH_FINDINGS.md",
                "output/doc/paper.docx",
            ],
            "user_stories": [
                {
                    "id": "US11",
                    "title": "Generate explainable advanced models",
                    "story": (
                        "As a policymaker, I want explainable advanced models and "
                        "state clusters so that the findings can support targeted "
                        "recommendations."
                    ),
                    "acceptance": [
                        "Nested cross-validation and advanced model outputs are generated.",
                        "SHAP values per state are exported for explainability.",
                        "Cluster and PCA outputs are saved for interpretation.",
                    ],
                },
                {
                    "id": "US12",
                    "title": "Prepare final project documentation",
                    "story": (
                        "As a student team, we want a paper-ready and report-ready "
                        "documentation set so that the project can be submitted and "
                        "presented cleanly."
                    ),
                    "acceptance": [
                        "The paper source and generated report documents are present.",
                        "Research findings are summarized in a reusable write-up.",
                        "Outputs are organized under data, figures, results, and output/doc.",
                    ],
                },
            ],
        },
    ]


def build_daily_scrum() -> list[dict]:
    return [
        {
            "day": "Sprint Day 1",
            "completed": "Defined the problem statement, research gap, project scope, and the main objectives for linking digital inclusion with circular economy behavior.",
            "next_step": "List the required official datasets and finalize the data collection plan.",
            "blockers": "No blocker. Initial scope was aligned with the research question.",
        },
        {
            "day": "Sprint Day 2",
            "completed": "Collected and verified raw inputs from TRAI, RBI, NITI Aayog, NPCI/Dataful, CPCB, Census, and NPCI monthly UPI files.",
            "next_step": "Standardize the source structure and prepare state-wise fields for integration.",
            "blockers": "State-level data came from multiple formats. Resolved by mapping each source to a common State field.",
        },
        {
            "day": "Sprint Day 3",
            "completed": "Built the data engineering pipeline and merged all verified sources using an inner join to retain only complete state records.",
            "next_step": "Engineer derived variables and export the unified dataset.",
            "blockers": "Not all states had complete coverage across every source. Resolved by restricting the main sample to 16 complete states.",
        },
        {
            "day": "Sprint Day 4",
            "completed": "Computed EWaste_PerCapita, Digital_Access_Score, CE_Behavior_Score, DCI, and descriptive statistics from the unified dataset.",
            "next_step": "Generate visual analysis and trend plots for interpretation.",
            "blockers": "Composite-weight justification needed documentation. Resolved by recording formulas and rationale in the paper and findings.",
        },
        {
            "day": "Sprint Day 5",
            "completed": "Generated EDA figures including the correlation heatmap, scatter plot, DCI ranking, distribution plots, and UPI trend visuals.",
            "next_step": "Run the main statistical hypothesis tests.",
            "blockers": "Figure readability needed refinement. Resolved by applying publication-style plotting settings.",
        },
        {
            "day": "Sprint Day 6",
            "completed": "Executed Pearson, Spearman, Kendall, OLS, multiple regression, mediation, threshold analysis, and diagnostic checks.",
            "next_step": "Interpret outputs and prepare robustness analysis for reviewer concerns.",
            "blockers": "Small sample size could weaken reviewer confidence. Mitigation planned through bootstrap and sensitivity analysis.",
        },
        {
            "day": "Sprint Day 7",
            "completed": "Ran reviewer-defense analyses including bootstrap confidence intervals, weight sensitivity, alternative mediation paths, and extended-sample checks.",
            "next_step": "Validate the findings through predictive modeling.",
            "blockers": "Expected criticism on weight subjectivity. Resolved by testing multiple alternative weighting schemes.",
        },
        {
            "day": "Sprint Day 8",
            "completed": "Compared seven machine-learning models using leave-one-out cross-validation and exported performance metrics and state-level predictions.",
            "next_step": "Run advanced explainability and clustering analyses.",
            "blockers": "Overfitting risk in a small dataset. Resolved by using leave-one-out validation instead of a simple train-test split.",
        },
        {
            "day": "Sprint Day 9",
            "completed": "Performed advanced ML tasks: XGBoost, nested CV, SHAP explainability, K-means clustering, repeated K-fold validation, and PCA visualization.",
            "next_step": "Convert analytical outputs into final paper and report material.",
            "blockers": "Model interpretation needed clearer justification. Resolved by exporting SHAP values and cluster summaries.",
        },
        {
            "day": "Sprint Day 10",
            "completed": "Compiled the research findings, final figures, results files, and paper-ready documentation.",
            "next_step": "Prepare project management documentation, backlog mapping, and submission-ready report sections.",
            "blockers": "Formatting iterations were required across figures and write-up. Resolved during documentation pass.",
        },
        {
            "day": "Sprint Day 11",
            "completed": "Mapped research objectives to epics, created user stories with acceptance criteria, and prepared the MS Planner-ready backlog and agile report.",
            "next_step": "Submit the final report and present the implementation trail during review or viva.",
            "blockers": "No blocker. Final documentation package completed.",
        },
    ]


def planner_rows(epics: list[dict]) -> list[dict]:
    rows = []
    for epic in epics:
        for story in epic["user_stories"]:
            rows.append(
                {
                    "Bucket": f"{epic['id']} - {epic['epic']}",
                    "Epic ID": epic["id"],
                    "Research Objective": epic["objective"],
                    "Task ID": story["id"],
                    "Task Title": story["title"],
                    "User Story": story["story"],
                    "Acceptance Criteria": " | ".join(story["acceptance"]),
                    "Priority": epic["priority"],
                    "Progress": epic["status"],
                    "Evidence": " | ".join(epic["evidence"]),
                }
            )
    return rows


def write_scrum_csv(rows: list[dict]) -> None:
    headers = ["Day", "Completed Update", "Next Step", "Blockers or Resolution"]
    with SCRUM_CSV_PATH.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=headers)
        writer.writeheader()
        for row in rows:
            writer.writerow(
                {
                    "Day": row["day"],
                    "Completed Update": row["completed"],
                    "Next Step": row["next_step"],
                    "Blockers or Resolution": row["blockers"],
                }
            )


def repo_summary() -> dict:
    return {
        "generated_on": date.today().isoformat(),
        "figure_count": count_files(BASE_DIR / "figures"),
        "result_count": count_files(BASE_DIR / "results"),
        "source_script_count": count_files(BASE_DIR / "src"),
    }


def write_csv(rows: list[dict]) -> None:
    headers = [
        "Bucket",
        "Epic ID",
        "Research Objective",
        "Task ID",
        "Task Title",
        "User Story",
        "Acceptance Criteria",
        "Priority",
        "Progress",
        "Evidence",
    ]
    with CSV_PATH.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=headers)
        writer.writeheader()
        writer.writerows(rows)


def build_markdown(epics: list[dict], summary: dict, scrum_rows: list[dict]) -> str:
    lines = []
    lines.append("# Minor Project Report: Product Backlog, User Stories, and Epic Implementation")
    lines.append("")
    lines.append("## 1. Project Context")
    lines.append(
        "This report converts the completed minor project on digital inclusion and circular economy "
        "behavior in India into an agile project-management view. Each research objective is treated "
        "as an epic, and the refined product backlog is expressed through user stories and acceptance criteria."
    )
    lines.append("")
    lines.append(
        f"- Report generated on: {summary['generated_on']}"
    )
    lines.append(
        f"- Repository evidence used: {summary['source_script_count']} source scripts, "
        f"{summary['result_count']} result files, {summary['figure_count']} figure files"
    )
    lines.append("- Project type: research-driven analytical pipeline, not a commercial software product")
    lines.append("")
    lines.append("## 2. Research Objectives Treated as Epics")
    lines.append("")
    lines.append("| Epic ID | Research Objective | Epic Name | Priority | Status |")
    lines.append("| --- | --- | --- | --- | --- |")
    for epic in epics:
        lines.append(
            f"| {epic['id']} | {epic['objective']} | {epic['epic']} | {epic['priority']} | {epic['status']} |"
        )
    lines.append("")
    lines.append("## 3. Refined Product Backlog")
    lines.append("")
    lines.append("| Epic ID | Backlog Item | Priority | Status | Implementation Summary |")
    lines.append("| --- | --- | --- | --- | --- |")
    for epic in epics:
        lines.append(
            f"| {epic['id']} | {epic['epic']} | {epic['priority']} | {epic['status']} | {epic['implementation']} |"
        )
    lines.append("")
    lines.append("## 4. User Stories and Acceptance Criteria")
    lines.append("")
    for epic in epics:
        lines.append(f"### {epic['id']}. {epic['epic']}")
        lines.append("")
        lines.append(f"**Research Objective:** {epic['objective']}")
        lines.append("")
        lines.append(f"**Implementation Performed:** {epic['implementation']}")
        lines.append("")
        for story in epic["user_stories"]:
            lines.append(f"**{story['id']} - {story['title']}**")
            lines.append("")
            lines.append(story["story"])
            lines.append("")
            lines.append("Acceptance Criteria:")
            for criterion in story["acceptance"]:
                lines.append(f"- {criterion}")
            lines.append("")
        lines.append("Implementation Evidence:")
        for item in epic["evidence"]:
            lines.append(f"- `{item}`")
        lines.append("")
    lines.append("### 4.1.b Daily Scrum (Daily call progress of the list of activities - standup updates)")
    lines.append("")
    lines.append(
        "The following daily scrum log summarizes how work progressed during the project in standup format."
    )
    lines.append("")
    lines.append("| Day | Completed Update | Next Step | Blockers / Resolution |")
    lines.append("| --- | --- | --- | --- |")
    for row in scrum_rows:
        lines.append(
            f"| {row['day']} | {row['completed']} | {row['next_step']} | {row['blockers']} |"
        )
    lines.append("")
    lines.append("## 5. Suggested MS Planner Structure")
    lines.append("")
    lines.append("- Create one bucket per epic: E1 to E6.")
    lines.append("- Create one task card per user story: US1 to US12.")
    lines.append("- Copy the acceptance criteria into the checklist field of each Planner task.")
    lines.append("- Mark all tasks and epics as `Completed`, because the implementation already exists in the repository.")
    lines.append("- Attach the listed evidence files or paste file paths in the task notes field.")
    lines.append("")
    lines.append("## 6. Implementation Summary")
    lines.append("")
    lines.append(
        "All defined epics were implemented in this project. The repository contains the complete trail of "
        "work: verified data files, engineered datasets, statistical outputs, machine-learning validation, "
        "advanced explainability, figures, and final paper/report documents."
    )
    lines.append("")
    lines.append("Key completed outputs:")
    lines.append("- Verified dataset inputs documented through official sources: TRAI, RBI, NITI Aayog, NPCI/Dataful, CPCB, and Census of India 2011")
    lines.append("- Statistical results: `results/ols_primary_results.txt`, `results/threshold_analysis.txt`, `results/mediation_analysis.txt`")
    lines.append("- Reviewer defense: `results/reviewer_defense_report.txt`")
    lines.append("- Predictive modeling outputs: `results/ml_model_comparison_digital_only.csv`, `results/ml_model_comparison_full.csv`")
    lines.append("- Advanced ML and explainability: `results/nested_cv_results.csv`, `results/shap_values_per_state.csv`, `results/state_clusters.csv`")
    lines.append("- Final documentation: `paper.tex`, `RESEARCH_FINDINGS.md`, `output/doc/paper.docx`")
    lines.append("")
    return "\n".join(lines)


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_table(document: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_cells = table.rows[0].cells
    for idx, header in enumerate(headers):
        set_cell_text(header_cells[idx], header, bold=True)
    set_repeat_table_header(table.rows[0])
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], value)


def apply_base_formatting(document: Document) -> None:
    for section in document.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)

    styles = document.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(11)
    styles["Title"].font.name = "Times New Roman"
    styles["Title"].font.size = Pt(18)
    styles["Heading 1"].font.name = "Times New Roman"
    styles["Heading 1"].font.size = Pt(14)
    styles["Heading 2"].font.name = "Times New Roman"
    styles["Heading 2"].font.size = Pt(12)
    styles["Heading 3"].font.name = "Times New Roman"
    styles["Heading 3"].font.size = Pt(11)


def add_bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.add_run(text)


def write_docx(epics: list[dict], summary: dict, scrum_rows: list[dict]) -> None:
    document = Document()
    apply_base_formatting(document)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Minor Project Report").bold = True
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Product Backlog, User Stories, and Epic Implementation")

    info = document.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info.add_run(f"Generated on {summary['generated_on']}")

    document.add_heading("1. Project Context", level=1)
    document.add_paragraph(
        "This report reframes the completed research project as an agile workflow. "
        "Each research objective is treated as an epic, and the completed work is "
        "expressed as a refined product backlog with user stories and acceptance criteria."
    )
    add_bullet(
        document,
        f"Repository evidence used: {summary['source_script_count']} source scripts, "
        f"{summary['result_count']} result files, {summary['figure_count']} figure files",
    )
    add_bullet(
        document,
        "Project type: research-driven analytical pipeline rather than a commercial software product.",
    )

    document.add_heading("2. Research Objectives Treated as Epics", level=1)
    add_table(
        document,
        ["Epic ID", "Research Objective", "Epic Name", "Priority", "Status"],
        [
            [
                epic["id"],
                epic["objective"],
                epic["epic"],
                epic["priority"],
                epic["status"],
            ]
            for epic in epics
        ],
    )

    document.add_paragraph()
    document.add_heading("3. Refined Product Backlog", level=1)
    add_table(
        document,
        ["Epic ID", "Backlog Item", "Priority", "Status", "Implementation Summary"],
        [
            [
                epic["id"],
                epic["epic"],
                epic["priority"],
                epic["status"],
                epic["implementation"],
            ]
            for epic in epics
        ],
    )

    document.add_paragraph()
    document.add_heading("4. User Stories and Acceptance Criteria", level=1)
    for epic in epics:
        document.add_heading(f"{epic['id']}. {epic['epic']}", level=2)
        p = document.add_paragraph()
        p.add_run("Research Objective: ").bold = True
        p.add_run(epic["objective"])

        p = document.add_paragraph()
        p.add_run("Implementation Performed: ").bold = True
        p.add_run(epic["implementation"])

        for story in epic["user_stories"]:
            document.add_heading(f"{story['id']} - {story['title']}", level=3)
            document.add_paragraph(story["story"])
            p = document.add_paragraph()
            p.add_run("Acceptance Criteria").bold = True
            for criterion in story["acceptance"]:
                add_bullet(document, criterion)

        p = document.add_paragraph()
        p.add_run("Implementation Evidence").bold = True
        for item in epic["evidence"]:
            add_bullet(document, item)

    document.add_heading(
        "4.1.b Daily Scrum (Daily call progress of the list of activities - standup updates)",
        level=2,
    )
    document.add_paragraph(
        "The following daily scrum log summarizes the standup-style progress updates recorded across the project workflow."
    )
    add_table(
        document,
        ["Day", "Completed Update", "Next Step", "Blockers / Resolution"],
        [
            [
                row["day"],
                row["completed"],
                row["next_step"],
                row["blockers"],
            ]
            for row in scrum_rows
        ],
    )

    document.add_heading("5. Suggested MS Planner Structure", level=1)
    add_bullet(document, "Create one bucket for each epic, from E1 through E6.")
    add_bullet(document, "Create one task card for each user story, from US1 through US12.")
    add_bullet(document, "Paste the acceptance criteria into the checklist field of each Planner task.")
    add_bullet(document, "Set task progress to Completed because the work is already implemented.")
    add_bullet(document, "Attach or reference evidence files in each task note.")

    document.add_heading("6. Implementation Summary", level=1)
    document.add_paragraph(
        "All defined epics were implemented in the current repository. The project includes "
        "verified source data, engineered datasets, statistical outputs, reviewer-defense checks, "
        "predictive modeling, advanced explainability, figures, and final documentation."
    )
    add_bullet(document, "Verified dataset inputs documented through official sources: TRAI, RBI, NITI Aayog, NPCI/Dataful, CPCB, and Census of India 2011")
    add_bullet(document, "Statistical outputs: results/ols_primary_results.txt, results/threshold_analysis.txt, results/mediation_analysis.txt")
    add_bullet(document, "Reviewer defense: results/reviewer_defense_report.txt")
    add_bullet(document, "Predictive modeling outputs: results/ml_model_comparison_digital_only.csv and results/ml_model_comparison_full.csv")
    add_bullet(document, "Advanced ML outputs: results/nested_cv_results.csv, results/shap_values_per_state.csv, results/state_clusters.csv")
    add_bullet(document, "Final documentation: paper.tex, RESEARCH_FINDINGS.md, output/doc/paper.docx")

    document.save(DOCX_PATH)


def main() -> None:
    epics = build_epics()
    scrum_rows = build_daily_scrum()
    summary = repo_summary()
    rows = planner_rows(epics)

    MD_PATH.write_text(build_markdown(epics, summary, scrum_rows), encoding="utf-8")
    write_csv(rows)
    write_scrum_csv(scrum_rows)
    write_docx(epics, summary, scrum_rows)

    print(f"Wrote {MD_PATH}")
    print(f"Wrote {CSV_PATH}")
    print(f"Wrote {SCRUM_CSV_PATH}")
    print(f"Wrote {DOCX_PATH}")


if __name__ == "__main__":
    main()
